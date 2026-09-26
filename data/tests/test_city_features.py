import json
import hashlib
from pathlib import Path
import pytest
from src.config.settings import PROCESSED
from src.contracts.models import CityFeature

def test_generated_city_features_validate_and_round_trip():
    assert (PROCESSED/"cities"/"all_city_features.json").exists(), "Run the full pipeline to create verified CityFeature records"
    records=json.loads((PROCESSED/"cities"/"all_city_features.json").read_text(encoding="utf-8"))
    assert len(records)==35
    for raw in records:
        city=CityFeature.model_validate(raw)
        assert city.city_id.isdigit() and len(city.city_id)==5
        assert city.features["population"].value>0
        assert city.features["population_density_per_km2"].value>0
        assert CityFeature.model_validate(city.model_dump(mode="json"))==city
        for measurement in city.features.values():
            assert measurement.quality=="missing" or measurement.provenance_ids
            assert set(measurement.provenance_ids)<={p.id for p in city.provenance}

def test_manifest_hashes_match_saved_raw_bytes():
    manifest=json.loads((PROCESSED/"manifests"/"data_manifest.json").read_text(encoding="utf-8"))
    assert manifest["datasets"]
    for entry in manifest["datasets"]:
        digest=entry["sha256"]
        assert len(digest)==64 and all(character in "0123456789abcdef" for character in digest.lower())
        if entry.get("retained", True) is False:
            continue
        raw_path=Path.cwd()/Path(entry["raw_path"])
        assert raw_path.is_file()
        assert hashlib.sha256(raw_path.read_bytes()).hexdigest()==digest

def test_new_climate_and_road_measurements_are_present_or_explicitly_missing():
    records=json.loads((PROCESSED/"cities"/"all_city_features.json").read_text(encoding="utf-8"))
    assert len(records)==35
    for raw in records:
        city=CityFeature.model_validate(raw)
        provenance_by_id={item.id:item for item in city.provenance}
        hot=city.features["hot_days_32c"]
        if hot.value is None:
            assert hot.missing_reason
        else:
            assert 0 <= hot.value <= 366
            hot_sources=[provenance_by_id[pid] for pid in hot.provenance_ids]
            assert 0 < len(hot_sources) <= 6 and len(hot_sources)%2==0
            assert all("within 100 km" in item.transformation for item in hot_sources if item.dataset_id=="noaa_normals_stations_1991-2020")
            assert sum(item.dataset_id.endswith("ge90f") for item in hot_sources) <= 3
        road_keys=("road_density_km_per_km2","intersection_density_per_km2","freeway_share","arterial_share","local_road_share")
        for key in road_keys:
            measurement=city.features[key]
            if measurement.value is None:
                assert measurement.missing_reason
            else:
                assert measurement.value >= 0
                assert measurement.provenance_ids
                assert all("EPSG:" in provenance_by_id[pid].transformation for pid in measurement.provenance_ids)
        aadt=city.features["average_aadt"]
        lane=city.features["lane_miles_per_km2"]
        assert aadt.value is None and lane.value is None
        assert aadt.missing_reason and lane.missing_reason

def test_excel_export_contains_all_city_values_and_measurement_details():
    from openpyxl import load_workbook
    workbook_path=PROCESSED/"cities"/"all_city_features.xlsx"
    assert workbook_path.is_file()
    workbook=load_workbook(workbook_path,read_only=True,data_only=True)
    assert workbook["City Values"].max_row==36
    headers=list(next(workbook["City Values"].iter_rows(values_only=True)))
    assert "hot_days_32c" in headers
    assert "average_aadt" in headers
    assert workbook["Measurement Details"].max_row==1+35*17
    workbook.close()

def test_missing_climate_is_explicit_and_not_zero():
    from src.datasets.noaa.process import missing_climate_features
    for measurement in missing_climate_features().values():
        assert measurement["value"] is None
        assert measurement["missing_reason"]
