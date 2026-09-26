"""Export the validated city release and source traceability to an Excel workbook."""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from src.config.settings import PROCESSED
from src.contracts.models import CityFeature


def export_city_features_excel(source: Path | None = None, output: Path | None = None) -> Path:
    """Write all configured city measurements, quality details, and provenance."""
    source = source or PROCESSED / "cities" / "all_city_features.json"
    output = output or PROCESSED / "cities" / "all_city_features.xlsx"
    raw_cities = json.loads(source.read_text(encoding="utf-8"))
    cities = [CityFeature.model_validate(item) for item in raw_cities]
    output.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    values_sheet = workbook.active
    values_sheet.title = "City Values"
    feature_keys = list(cities[0].features)
    identity = ["city_id", "display_name", "official_name", "geography_type", "geography_vintage", "state_codes"]
    values_sheet.append(identity + feature_keys)
    details_sheet = workbook.create_sheet("Measurement Details")
    details_sheet.append(["city_id", "city_name", "feature", "value", "unit", "quality", "missing_reason", "provenance_ids"])
    provenance_sheet = workbook.create_sheet("Provenance")
    provenance_sheet.append([
        "city_id", "city_name", "provenance_id", "source_name", "dataset_id", "period", "source_url",
        "retrieved_at", "source_geography", "target_geography", "transformation", "assumptions", "raw_sha256",
    ])

    for city in cities:
        values_sheet.append([
            city.city_id, city.display_name, city.official_name, city.geography_type,
            city.geography_vintage, ", ".join(city.state_codes),
            *[city.features[key].value for key in feature_keys],
        ])
        for key in feature_keys:
            measurement = city.features[key]
            details_sheet.append([
                city.city_id, city.display_name, key, measurement.value, measurement.unit,
                measurement.quality, measurement.missing_reason, ", ".join(measurement.provenance_ids),
            ])
        for item in city.provenance:
            provenance_sheet.append([
                city.city_id, city.display_name, item.id, item.source_name, item.dataset_id, item.period,
                item.source_url, item.retrieved_at, item.source_geography, item.target_geography,
                item.transformation, " | ".join(item.assumptions), item.raw_sha256,
            ])

    notes = workbook.create_sheet("Notes")
    notes.append(["Topic", "Definition / note"])
    notes.append(["Geography", "Every record is an official CBSA using the frozen 2024 Census CBSA and county membership."])
    notes.append(["hot_days_32c", "1991–2020 NOAA annual normal: mean annual days when daily Tmax is at least 90°F (32.222…°C), averaged over up to three qualifying stations within 100 km."])
    notes.append(["Road sources", "2024 Census TIGER/Line county EDGES for counties in the 35 configured CBSAs. Raw ZIPs were removed after processing; their SHA-256 values remain in provenance and the data manifest."])
    notes.append(["Road density", "Unique qualifying drivable road centerline kilometers divided by CBSA county land area in km²."])
    notes.append(["Intersections", "Census TNID endpoints with at least three distinct incident edges, nearby nodes merged within 20 m, and CBSA-boundary nodes excluded within 5 m."])
    notes.append(["Functional class", "Freeway proxy=S1100+S1630; arterial proxy=S1200; local/residual=S1400+S1640+S1730, weighted by road length."])
    notes.append(["AADT and lane density", "Missing for all metros because no current complete comparable HPMS segment layer was obtained. Blank cells mean missing; they are not zero."])
    notes.append(["Use limitation", "Road metrics are public road-environment proxies, not autonomous-driving safety, readiness, or deployment-approval measures."])

    header_fill = PatternFill("solid", fgColor="17324D")
    header_font = Font(color="FFFFFF", bold=True)
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 30
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for column in sheet.columns:
            letter = get_column_letter(column[0].column)
            max_length = max((len(str(cell.value)) for cell in column if cell.value is not None), default=8)
            sheet.column_dimensions[letter].width = min(max(max_length + 2, 12), 56)

    for sheet, name in ((values_sheet, "CityFeatureValues"), (details_sheet, "MeasurementDetails"), (provenance_sheet, "CityProvenance")):
        table = Table(displayName=name, ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
        sheet.add_table(table)
    for row in values_sheet.iter_rows(min_row=2, min_col=len(identity) + 1):
        for cell in row:
            cell.number_format = "0.000"
    for row in details_sheet.iter_rows(min_row=2, min_col=4, max_col=4):
        row[0].number_format = "0.000"

    workbook.save(output)
    manifest_path = PROCESSED / "manifests" / "data_manifest.json"
    if manifest_path.exists():
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        relative_output = output.resolve().relative_to(Path.cwd().resolve()).as_posix()
        outputs = manifest.setdefault("processed_outputs", [])
        if relative_output not in outputs:
            outputs.append(relative_output)
            manifest_path.write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    return output


if __name__ == "__main__":
    print(export_city_features_excel())
